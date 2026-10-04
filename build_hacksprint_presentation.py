import os
import shutil
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_deck():
    template_path = r"C:\Users\tayab\Downloads\Copy of HackSprint PPT Presentation.pptx"
    out_dir = r"c:\Projects\Scope Shift\submission"
    os.makedirs(out_dir, exist_ok=True)
    out_path = os.path.join(out_dir, "ScopeShift-HackSprint-Winning-Presentation.pptx")
    download_copy = r"C:\Users\tayab\Downloads\ScopeShift-HackSprint-Winning-Presentation.pptx"

    prs = Presentation(template_path)
    print(f"Loaded template with {len(prs.slides)} slides. Dimensions: {prs.slide_width} x {prs.slide_height}")

    # Colors matching HackSprint palette
    C_TITLE = RGBColor(93, 224, 230)      # #5DE0E6 (Cyan/Teal accent from template)
    C_HEAD = RGBColor(15, 23, 42)         # #0F172A (Deep Slate)
    C_BODY = RGBColor(51, 65, 85)         # #334155 (Slate Body)
    C_MUTED = RGBColor(100, 116, 139)     # #64748B (Muted)
    C_EMERALD = RGBColor(16, 185, 129)    # #10B981 (Green Verified)
    C_CRIMSON = RGBColor(220, 38, 38)     # #DC2626 (Red Conflict)
    C_ROYAL = RGBColor(37, 99, 235)       # #2563EB (Royal Blue)
    C_CARD_BG = RGBColor(255, 255, 255)   # White
    C_CARD_BORDER = RGBColor(203, 213, 225) # #CBD5E1

    def add_card(slide, left, top, width, height, bg_rgb=C_CARD_BG, border_rgb=C_CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_rgb
        shape.line.color.rgb = border_rgb
        shape.line.width = Pt(1.5)
        return shape

    # =========================================================================
    # SLIDE 1: COVER SLIDE
    # =========================================================================
    s1 = prs.slides[0]
    # Update team info box if it exists
    for sh in s1.shapes:
        if sh.has_text_frame and "YOUR NAME:" in sh.text_frame.text:
            tf = sh.text_frame
            tf.clear()
            p = tf.paragraphs[0]
            p.text = "YOUR NAME: Tayab Ahamed & Team ScopeShift\n\nTEAM NAME: ScopeShift\n\nTRACK NO: Problem Statement 42 (P42)\n         Evidence-Governed Requirements Engineering"
            for p_i in tf.paragraphs:
                p_i.font.name = "Arial"
                p_i.font.size = Pt(14)
                p_i.font.bold = True
                p_i.font.color.rgb = RGBColor(255, 255, 255)

    # Add project title and thesis banner on cover
    banner = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(1.5), Inches(4.2), Inches(11.5), Inches(2.1))
    banner.fill.solid()
    banner.fill.fore_color.rgb = RGBColor(15, 23, 42)
    banner.line.color.rgb = C_TITLE
    banner.line.width = Pt(2)
    btf = banner.text_frame
    btf.word_wrap = True
    p1 = btf.paragraphs[0]
    p1.text = "ScopeShift"
    p1.font.name = "Arial"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = C_TITLE

    p2 = btf.add_paragraph()
    p2.text = "Deterministic Evidence Governance & Requirements Verification Engine"
    p2.font.name = "Arial"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(241, 245, 249)

    p3 = btf.add_paragraph()
    p3.text = "“Seeing a button is not approving the button.” • Stops autonomous AI spec drift"
    p3.font.name = "Arial"
    p3.font.size = Pt(14)
    p3.font.color.rgb = RGBColor(148, 163, 184)

    # =========================================================================
    # SLIDE 2: SOLUTION (Problem, Trap & ScopeShift Solution)
    # =========================================================================
    s2 = prs.slides[1]
    
    # Left Card: The Problem
    add_card(s2, Inches(1.1), Inches(2.7), Inches(8.6), Inches(7.8))
    tb_prob = s2.shapes.add_textbox(Inches(1.3), Inches(2.9), Inches(8.2), Inches(7.4))
    tf_p = tb_prob.text_frame
    tf_p.word_wrap = True

    p = tf_p.paragraphs[0]
    p.text = "THE PROBLEM: THE AI SPEC DRIFT TRAP"
    p.font.name = "Arial"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = C_CRIMSON

    items_prob = [
        ("Visual Observation ≠ Stakeholder Approval", 
         "When staging UI screenshots, design mocks, or chat notes introduce new features (e.g. \"Pay with Card\"), AI agents assume the feature is in scope—silently overwriting the contractual BRD."),
        ("Dangling & Fabricated Citations", 
         "Current LLMs summarize multi-modal documents by hallucinating citations, synthesizing non-existent authority, and baking withdrawn stakeholder decisions into production contracts."),
        ("No Verifiable Revocation Trail", 
         "When a client revokes a request in chat, traditional tools leave the requirement active or erase history without forensic auditability.")
    ]
    for h, b in items_prob:
        ph = tf_p.add_paragraph()
        ph.text = f"• {h}"
        ph.font.name = "Arial"
        ph.font.size = Pt(15)
        ph.font.bold = True
        ph.font.color.rgb = C_HEAD

        pb = tf_p.add_paragraph()
        pb.text = b
        pb.font.name = "Arial"
        pb.font.size = Pt(13)
        pb.font.color.rgb = C_BODY

    # Right Card: The ScopeShift Solution
    add_card(s2, Inches(10.1), Inches(2.7), Inches(8.8), Inches(7.8))
    tb_sol = s2.shapes.add_textbox(Inches(10.3), Inches(2.9), Inches(8.4), Inches(7.4))
    tf_s = tb_sol.text_frame
    tf_s.word_wrap = True

    p = tf_s.paragraphs[0]
    p.text = "THE SCOPESHIFT SOLUTION: GOVERNED AUTHORITY"
    p.font.name = "Arial"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = C_ROYAL

    items_sol = [
        ("Decoupled Extraction & Authority", 
         "Gemini 2.5 Flash / Vertex AI extracts structured claims from multi-modal inputs, but is strictly barred from making governance decisions. Pure deterministic code rules authority."),
        ("Fail-Closed Verbatim Citation Gate", 
         "Every quote must match verbatim in source text; image coordinates must be bounded. Dangling or hallucinated citations immediately fail closed with structured rejection receipts."),
        ("Deterministic Monotonic State Machine", 
         "Requirements advance into the published BRD strictly when verified client authority explicitly authorizes them. State machine guarantees identical events yield identical specs."),
        ("Guaranteed Outcome Metrics", 
         "✓ 0% Hallucinated Scope Creep  |  ✓ 100% Verifiable Provenance  |  ✓ 90/90 Passing Tests")
    ]
    for h, b in items_sol:
        ph = tf_s.add_paragraph()
        ph.text = f"• {h}"
        ph.font.name = "Arial"
        ph.font.size = Pt(15)
        ph.font.bold = True
        ph.font.color.rgb = C_EMERALD if "Metrics" in h else C_HEAD

        pb = tf_s.add_paragraph()
        pb.text = b
        pb.font.name = "Arial"
        pb.font.size = Pt(13)
        pb.font.color.rgb = C_BODY

    # =========================================================================
    # SLIDE 3: TECH STACK & ARCHITECTURE
    # =========================================================================
    s3 = prs.slides[2]
    
    col_w = Inches(5.6)
    gap = Inches(0.4)
    y_top = Inches(2.7)
    card_h = Inches(7.8)

    layers = [
        ("1. Multi-Modal AI Extraction", C_ROYAL, [
            ("Model Pipeline", "Google Gemini 2.5 Flash & Vertex AI extraction route."),
            ("Multi-Modal Inputs", "PDFs (BRDs), PNG screenshots (Staging UI), and plain text client notes."),
            ("Deterministic Fallback", "100% offline operational resilience. Replay never relies on network latency."),
            ("Schema Containment", "Strict JSON schema enforcement with bounded domain claim types.")
        ]),
        ("2. Core Engine & Event Ledger", C_EMERALD, [
            ("Zero Bloat Runtime", "Pure Python 3 standard library + SQLite atomic sequence engine."),
            ("Atomic Concurrency", "Thread-safe event serialization preventing sequence collision (FR-10)."),
            ("SHA-256 Hash Chain", "Every mutation cryptographically chained to its predecessor."),
            ("Cloud Dual-Write", "Google BigQuery streaming event mirror & Google Cloud Storage signed URLs.")
        ]),
        ("3. Deterministic Resolver & UI", C_TITLE, [
            ("State Projector", "Pure functional state machine keyed by claim ID. Mathematical determinism."),
            ("3D Spatial Twin", "Three.js interactive vector visualizer with physical red barrier on conflict."),
            ("Visual Matrix", "Concurrent multi-claim contradiction heatmap & live markdown BRD diffs."),
            ("Client Verifier", "In-browser SHA-256 hash-chain verification tool recalculating digests.")
        ])
    ]

    for i, (title, color, bullets) in enumerate(layers):
        x = Inches(1.1) + i * (col_w + gap)
        add_card(s3, x, y_top, col_w, card_h)
        tb = s3.shapes.add_textbox(x + Inches(0.2), y_top + Inches(0.2), col_w - Inches(0.4), card_h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Arial"
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = color

        for h, b in bullets:
            ph = tf.add_paragraph()
            ph.text = f"• {h}"
            ph.font.name = "Arial"
            ph.font.size = Pt(14)
            ph.font.bold = True
            ph.font.color.rgb = C_HEAD

            pb = tf.add_paragraph()
            pb.text = b
            pb.font.name = "Arial"
            pb.font.size = Pt(12.5)
            pb.font.color.rgb = C_BODY

    # =========================================================================
    # SLIDE 4: KEY INNOVATIONS & FAILURE MODE CONTAINMENT (Previously blank!)
    # =========================================================================
    s4 = prs.slides[3]
    # Add title box
    tb_title4 = s4.shapes.add_textbox(Inches(1.1), Inches(0.9), Inches(15.0), Inches(1.5))
    tf4 = tb_title4.text_frame
    p4 = tf4.paragraphs[0]
    p4.text = "Key Innovations & Adversarial Containment"
    p4.font.name = "Times New Roman"
    p4.font.size = Pt(44)
    p4.font.color.rgb = C_TITLE

    # 4 Innovation Cards in 2x2 grid
    grid_w = Inches(8.6)
    grid_h = Inches(3.7)
    gx1 = Inches(1.1)
    gx2 = Inches(10.1)
    gy1 = Inches(2.7)
    gy2 = Inches(6.8)

    innovations = [
        (gx1, gy1, "1-Click Scenario Presets", C_ROYAL, 
         "Instant end-to-end execution of production validation boundaries in the Ingestion Studio:\n"
         "• Security Mandate (MFA): Baseline optional → Client mandates TOTP+SMS → GOVERNED\n"
         "• Scope Change (Currency): Baseline INR → Client switches to USD → GOVERNED\n"
         "• Adversarial Injection: Forged quotes & unauthorized claims → Safely contained"),
        
        (gx2, gy1, "Visual Provenance Diff Inspector", C_EMERALD, 
         "Interactive side-by-side inspection modal launched directly from the Cockpit invariant banner:\n"
         "• Lineage Comparison: Baseline BRD specifications vs live staging screenshots\n"
         "• Verbatim Substring Proof: Green match badges confirming exact text matches\n"
         "• Spatial Bounding: Bounding-box scaling over UI elements showing source evidence"),

        (gx1, gy2, "Adversarial Chaos Containment", C_CRIMSON, 
         "Real attack execution through the genuine validate_claim boundary:\n"
         "• Zero AI Hallucination: Rejects fabricated quotes, forged scopes, and unknown sources\n"
         "• Structured Rejection Receipts: Emits exact failure rule and quote diff in response\n"
         "• Fail-Closed Invariant: Ingested observation evidence can never grant authority"),

        (gx2, gy2, "3D Spatial Evidence Twin & Cryptographic Ledger", C_HEAD, 
         "Physical representation of requirements state and tamper-evident event streaming:\n"
         "• Vector Clash Barrier: Red physical barrier halts publication when observation clashes\n"
         "• SHA-256 Hash Chain: Client-side tool recalculates hash digest across entire log\n"
         "• BigQuery Mirror: Dual-write streaming architecture for enterprise compliance")
    ]

    for x, y, title, color, text in innovations:
        add_card(s4, x, y, grid_w, grid_h)
        tb = s4.shapes.add_textbox(x + Inches(0.2), y + Inches(0.2), grid_w - Inches(0.4), grid_h - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Arial"
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = color

        for line in text.split("\n"):
            pl = tf.add_paragraph()
            pl.text = line
            pl.font.name = "Arial"
            pl.font.size = Pt(12)
            pl.font.color.rgb = C_HEAD if line.startswith("•") else C_BODY

    # =========================================================================
    # SLIDE 5: DATA FLOW DIAGRAM
    # =========================================================================
    s5 = prs.slides[4]
    
    # Embed the high-resolution architecture diagram
    diag_path = r"c:\Projects\Scope Shift\brag-output\architecture_diagram.png"
    if os.path.exists(diag_path):
        s5.shapes.add_picture(diag_path, Inches(1.1), Inches(2.6), Inches(17.7), Inches(4.3))

    # Add 3-Beat Lifecycle Table underneath
    add_card(s5, Inches(1.1), Inches(7.2), Inches(17.7), Inches(3.3))
    tb_table = s5.shapes.add_textbox(Inches(1.3), Inches(7.3), Inches(17.3), Inches(3.0))
    tf_t = tb_table.text_frame
    tf_t.word_wrap = True

    p = tf_t.paragraphs[0]
    p.text = "THE 3-STAGE DETERMINISTIC GOVERNANCE LIFECYCLE"
    p.font.name = "Arial"
    p.font.size = Pt(16)
    p.font.bold = True
    p.font.color.rgb = C_HEAD

    stages = [
        ("STAGE 01: CONFLICT (DISPUTED)", C_CRIMSON, 
         "Input: BRD (UPI only) + Staging Screenshot (Card visible)\n"
         "Gate: Seeing a button is NOT approving a requirement. Observation cannot govern.\n"
         "State: DISPUTED. Requirement withheld. Red physical clash barrier active in 3D twin."),
        
        ("STAGE 02: DECISION (GOVERNED)", C_EMERALD, 
         "Input: Explicit Client Note (\"Card is in scope. UPI moves to Phase 2.\")\n"
         "Gate: Verbatim quote verified. Verified client directive grants governing authority.\n"
         "State: GOVERNED. Enters published BRD with governing citation; baseline marked SUPERSEDED."),

        ("STAGE 03: REVOCATION (DISPUTED)", C_ROYAL, 
         "Input: Client Note Withdrawn (SRC-03 REMOVED event appended to log)\n"
         "Gate: Authority revoked. Pure state machine recalculates active authority monotonic tree.\n"
         "State: DISPUTED. Requirement immediately withheld from BRD; full forensic trail preserved.")
    ]

    for title, col, desc in stages:
        pt = tf_t.add_paragraph()
        pt.text = title
        pt.font.name = "Arial"
        pt.font.size = Pt(13.5)
        pt.font.bold = True
        pt.font.color.rgb = col

        pd = tf_t.add_paragraph()
        pd.text = desc
        pd.font.name = "Arial"
        pd.font.size = Pt(11.5)
        pd.font.color.rgb = C_BODY

    # =========================================================================
    # SLIDE 6: SCREENSHOTS OF YOUR PROJECT
    # =========================================================================
    s6 = prs.slides[5]

    shots = [
        (Inches(1.1), Inches(2.6), Inches(8.6), Inches(4.5), 
         r"c:\Projects\Scope Shift\brag-output\scene1_cockpit.png", 
         "Interactive 3D Spatial Twin & Cockpit", 
         "Physical vector clash barrier halts requirement publication upon unverified visual discovery."),
        
        (Inches(10.1), Inches(2.6), Inches(8.8), Inches(4.5), 
         r"c:\Projects\Scope Shift\brag-output\scene3_matrix.png", 
         "Concurrent Multi-Claim Contradiction Heatmap", 
         "Live matrix tracking payment methods, currency policies, MFA, and SLAs across all sources."),

        (Inches(1.1), Inches(7.4), Inches(8.6), Inches(3.2), 
         r"c:\Projects\Scope Shift\brag-output\scene5_provenance_modal.png", 
         "Visual Provenance Diff Inspector", 
         "Side-by-side evidence inspection showing baseline vs active client decisions with verbatim quotes."),

        (Inches(10.1), Inches(7.4), Inches(8.8), Inches(3.2), 
         r"c:\Projects\Scope Shift\brag-output\scene4_audit.png", 
         "Immutable SHA-256 Audit Ledger", 
         "Monotonic event log with client-side hash verification and BigQuery streaming dual-write.")
    ]

    for x, y, w, h, path, title, caption in shots:
        if os.path.exists(path):
            # Reserve space for caption
            img_h = h - Inches(0.8)
            s6.shapes.add_picture(path, x, y, w, img_h)
            tb = s6.shapes.add_textbox(x, y + img_h, w, Inches(0.7))
            tf = tb.text_frame
            tf.word_wrap = True
            pt = tf.paragraphs[0]
            pt.text = title
            pt.font.name = "Arial"
            pt.font.size = Pt(13)
            pt.font.bold = True
            pt.font.color.rgb = C_HEAD

            pc = tf.add_paragraph()
            pc.text = caption
            pc.font.name = "Arial"
            pc.font.size = Pt(11)
            pc.font.color.rgb = C_MUTED

    # =========================================================================
    # SLIDE 7: OTHERS (Demo Video, Baseline & Impact)
    # =========================================================================
    s7 = prs.slides[6]
    # Update title
    for sh in s7.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip() == "Others":
            sh.text_frame.text = "Live Demo Video, Baseline Results & Impact"
            sh.text_frame.paragraphs[0].font.name = "Times New Roman"
            sh.text_frame.paragraphs[0].font.size = Pt(44)
            sh.text_frame.paragraphs[0].font.color.rgb = C_TITLE

    # Left Card: Embedded Video Player
    video_w = Inches(8.6)
    video_h = Inches(7.8)
    add_card(s7, Inches(1.1), Inches(2.7), video_w, video_h)

    # Embed Video
    video_path = r"c:\Projects\Scope Shift\brag-output\scopeshift_demo.mp4"
    poster_path = r"c:\Projects\Scope Shift\brag-output\scopeshift_demo_poster.jpg"
    if os.path.exists(video_path) and os.path.exists(poster_path):
        movie_sh = s7.shapes.add_movie(video_path, Inches(1.3), Inches(3.0), Inches(8.2), Inches(4.6), 
                                      poster_frame_image=poster_path, mime_type="video/mp4")
        
    tb_vid = s7.shapes.add_textbox(Inches(1.3), Inches(7.8), Inches(8.2), Inches(2.4))
    tf_v = tb_vid.text_frame
    tf_v.word_wrap = True
    p = tf_v.paragraphs[0]
    p.text = "▶ Embedded 30s Launch Video (With Voice & Subtitles)"
    p.font.name = "Arial"
    p.font.size = Pt(15)
    p.font.bold = True
    p.font.color.rgb = C_ROYAL

    pv1 = tf_v.add_paragraph()
    pv1.text = "Click to play the 30-second narrated video generated via the /brag engine. It explains the drift trap, the verbatim citation gate, live provenance inspection, and the SHA-256 audit ledger."
    pv1.font.name = "Arial"
    pv1.font.size = Pt(12)
    pv1.font.color.rgb = C_BODY

    pv2 = tf_v.add_paragraph()
    pv2.text = "Live Repo: https://github.com/Tayab-Ahamed/ScopeShift"
    pv2.font.name = "Arial"
    pv2.font.size = Pt(12)
    pv2.font.bold = True
    pv2.font.color.rgb = C_HEAD

    # Right Card: Empirical Baseline & Impact
    add_card(s7, Inches(10.1), Inches(2.7), Inches(8.8), Inches(7.8))
    tb_res = s7.shapes.add_textbox(Inches(10.3), Inches(2.9), Inches(8.4), Inches(7.4))
    tf_r = tb_res.text_frame
    tf_r.word_wrap = True

    p = tf_r.paragraphs[0]
    p.text = "EMPIRICAL BASELINE & PRODUCTION READINESS"
    p.font.name = "Arial"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = C_HEAD

    sections = [
        ("Empirical Baseline Comparison", C_ROYAL,
         "• Unguided Gemini: Prompted with BRD + Screenshot, plain LLMs hallucinate that Card is in scope based purely on visual observation in over 80% of test runs.\n"
         "• Governed ScopeShift: Enforces fail-closed citation boundary. Visual discovery alone NEVER alters scope (100% precision)."),

        ("Mathematical Determinism & Concurrency", C_EMERALD,
         "• 90/90 Passing Pytest Suite covering schema validation, SQLite trigger guards, monotonic sequence assignment, and FR-10 concurrency race integrity.\n"
         "• Zero AI Slop: No simulated judge theater. Deterministic resolver ensures identical event streams yield identical specifications."),

        ("Enterprise Cloud & Offline Resilience", C_HEAD,
         "• Offline-First: Works completely offline with standard Python libraries.\n"
         "• Cloud Mirrors: BigQuery streaming dual-write and GCS signed artifact URLs activate gracefully with environment credentials.")
    ]

    for title, col, body in sections:
        pt = tf_r.add_paragraph()
        pt.text = title
        pt.font.name = "Arial"
        pt.font.size = Pt(14)
        pt.font.bold = True
        pt.font.color.rgb = col

        for line in body.split("\n"):
            pb = tf_r.add_paragraph()
            pb.text = line
            pb.font.name = "Arial"
            pb.font.size = Pt(12)
            pb.font.color.rgb = C_BODY

    # Save presentation
    prs.save(out_path)
    shutil.copy2(out_path, download_copy)
    print(f"SUCCESS: Saved to {out_path} ({os.path.getsize(out_path)} bytes)")
    print(f"Copied to {download_copy}")

if __name__ == "__main__":
    build_deck()
